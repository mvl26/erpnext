# Copyright (c) 2026, Công ty TNHH Miyano Việt Nam

"""Xuất báo cáo Đối chiếu mua vào - bán ra ra xlsx theo bố cục mẫu `Công nợ T10_Himedic_20.10.xlsx`.

Dòng 1 tiêu đề, dòng 2 ghi chú cơ sở kỳ, dòng 3 các khối KHÁCH HÀNG / NHÀ CUNG CẤP /
ĐỐI CHIẾU, dòng 4 tiêu đề cột, dữ liệu từ dòng 5, rồi TỔNG CỘNG và phần tóm tắt. Các
mảnh tách từ một dòng bán gộp ô bên trái (STT, số HĐ, tên hàng...) như file mẫu. Nút
Excel mặc định của ERP vẫn còn, nhưng ra bảng phẳng không có các khối.
"""

import json
from io import BytesIO

import frappe
from frappe import _
from frappe.utils import flt, getdate
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from erpnext.accounts.report.doi_chieu_mua_vao_ban_ra.doi_chieu_mua_vao_ban_ra import (
	COT_DOI_CHIEU,
	execute,
	khoang_ky,
)
from erpnext.accounts.report.doi_chieu_mua_vao_ban_ra.phan_bo import KHOP

REPORT = "Doi Chieu Mua Vao Ban Ra"
GHI_CHU_KY = "Kỳ lọc theo ngày hóa đơn đầu ra; hóa đơn đầu vào ghép được vẫn hiện dù ngày ở kỳ khác."
_TIEN_AM_DO = "#,##0;[Red](#,##0)"

_TIEU_DE_KHOI = [
	"Số HĐ",
	"Ngày HĐ",
	"Tên hàng hóa (nguyên văn hóa đơn)",
	"ĐVT",
	"SL",
	"Đơn giá (chưa VAT)",
	"Thành tiền (chưa VAT)",
	"Thuế suất",
	"Tiền thuế",
	"Tổng sau thuế",
]
_TRUONG_KHOI = [
	"so_hd",
	"ngay_hd",
	"ten_hang",
	"dvt",
	"sl",
	"don_gia",
	"thanh_tien",
	"thue_suat",
	"tien_thue",
	"tong",
]
_DINH_DANG = {
	"ngay_hd": "dd/mm/yyyy",
	"sl": "#,##0.##",
	"don_gia": "#,##0.##",
	"thanh_tien": _TIEN_AM_DO,
	"thue_suat": "0%",
	"tien_thue": _TIEN_AM_DO,
	"tong": _TIEN_AM_DO,
	"ty_le_gia": "0.000000",
	"chenh_lech_pt": "0.00%",
	"chenh_lech_tien": _TIEN_AM_DO,
}
_RONG = {"ten_hang": 50, "so_hd": 13, "ngay_hd": 11, "dvt": 8, "sl": 7, "thue_suat": 9, "ghi_chu": 45}
# Cột giống nhau giữa các mảnh của một dòng bán: gộp ô theo chiều dọc.
_GOP = {"stt", "kh", "so_hd_ban", "ngay_hd_ban", "ten_hang_ban", "dvt_ban", "don_gia_ban", "thue_suat_ban"}

# Màu và chữ lấy đúng từ file mẫu `Công nợ T10_Himedic_20.10.xlsx`; trang web dùng cùng mã màu.
_FONT = "Times New Roman"
_CO_CHU = 11  # mọi ô trừ tiêu đề dòng 1
_MAU_STT_NCC = "1F2A4A"  # ô STT và khối NHÀ CUNG CẤP (dòng 3-4)
_MAU_KH = "5B8DD9"  # khối KHÁCH HÀNG (dòng 3-4)
_MAU_DOI_CHIEU = "595959"  # khối ĐỐI CHIẾU — file mẫu không có, chọn xám để tách khỏi hai khối gốc
_MAU_SO_HD_RA = "FFF2CC"  # cột Số HĐ đầu ra
_MAU_TONG = "E8EEF8"  # dòng TỔNG CỘNG
_MAU_VIEN = "BFBFBF"
_DO = Font(name=_FONT, size=_CO_CHU, color="C00000")


def _nen(mau):
	return PatternFill("solid", fgColor=mau)


def _kiem_quyen(filters):
	if not frappe.get_doc("Report", REPORT).is_permitted():
		frappe.throw(_("Bạn không có quyền xem báo cáo này."), frappe.PermissionError)
	return frappe._dict(json.loads(filters) if isinstance(filters, str) else filters)


def _tieu_de(filters):
	"""(tên khách in hoa, tiêu đề dòng 1) — trang web và file Excel dùng chung."""
	_tu, _den, nhan = khoang_ky(filters)
	ten_kh = frappe.db.get_value("Customer", filters.customer, "customer_name") or filters.customer or ""
	return ten_kh.upper(), f"TỔNG HỢP ĐẦU VÀO - RA CỦA KHÁCH HÀNG {ten_kh.upper()} {nhan}"


@frappe.whitelist()
def du_lieu_trang(filters):
	"""Dữ liệu cho trang `doi-chieu-mua-ban`: cùng `execute()` với file Excel nên hai nơi luôn khớp."""
	filters = _kiem_quyen(filters)
	if not (filters.get("company") and filters.get("customer")):
		return {"ten_kh": "", "tieu_de": "", "data": []}
	ten_kh, tieu_de = _tieu_de(filters)
	return {"ten_kh": ten_kh, "tieu_de": tieu_de, "ghi_chu_ky": GHI_CHU_KY, "data": execute(filters)[1]}


@frappe.whitelist()
def xuat(filters):
	filters = _kiem_quyen(filters)
	tu, den, _nhan = khoang_ky(filters)
	frappe.response.filename = f"Doi chieu mua ban_{filters.customer}_{tu:%d%m%Y}-{den:%d%m%Y}.xlsx"
	frappe.response.filecontent = dung_workbook(filters)
	frappe.response.type = "binary"


def dung_workbook(filters):
	filters = frappe._dict(filters)
	_cot, data = execute(filters)
	ten_kh, tieu_de = _tieu_de(filters)

	wb = Workbook()
	ws = wb.active
	ws.title = "Mua vào - Bán ra"
	chu = Font(name=_FONT, size=_CO_CHU)
	chu_tong = Font(name=_FONT, size=_CO_CHU, bold=True)
	chu_tieu_de_cot = Font(name=_FONT, size=_CO_CHU, bold=True, color="FFFFFF")
	vien = Border(*(Side(style="thin", color=_MAU_VIEN),) * 4)
	giua = Alignment(horizontal="center", vertical="center", wrap_text=True)
	tren = Alignment(vertical="top", wrap_text=True)

	truong = [("stt", None), ("kh", None)]
	truong += [(f"{t}_ban", t) for t in _TRUONG_KHOI]
	truong += [("ncc", None)] + [(f"{t}_mua", t) for t in _TRUONG_KHOI]
	truong += [(f, f) for _nhan, f, _kieu, _rong in COT_DOI_CHIEU]
	so_cot = len(truong)

	ws.cell(1, 1, tieu_de).font = Font(name=_FONT, size=18)
	ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=so_cot)
	ws.cell(1, 1).alignment = giua
	ws.row_dimensions[1].height = 23.25
	ws.cell(2, 1, GHI_CHU_KY).font = Font(name=_FONT, size=_CO_CHU, italic=True)

	khoi_ban, khoi_mua = 2, 2 + 1 + len(_TRUONG_KHOI)
	khoi_dc = khoi_mua + 1 + len(_TRUONG_KHOI)

	def mau_cot(c):
		if c == 1:
			return _MAU_STT_NCC
		if c < khoi_mua:
			return _MAU_KH
		return _MAU_STT_NCC if c < khoi_dc else _MAU_DOI_CHIEU

	for cot_dau, cot_cuoi, nhan_khoi in (
		(khoi_ban, khoi_mua - 1, f"KHÁCH HÀNG {ten_kh}"),
		(khoi_mua, khoi_dc - 1, "NHÀ CUNG CẤP"),
		(khoi_dc, so_cot, "ĐỐI CHIẾU"),
	):
		ws.cell(3, cot_dau, nhan_khoi)
		ws.merge_cells(start_row=3, start_column=cot_dau, end_row=3, end_column=cot_cuoi)
	for c in range(1, so_cot + 1):
		o = ws.cell(3, c)
		o.font, o.alignment, o.fill = chu_tieu_de_cot, giua, _nen(mau_cot(c))

	tieu_de = ["STT", "KH", "Số HĐ đầu ra", *_TIEU_DE_KHOI[1:], "NCC", "Số HĐ đầu vào", *_TIEU_DE_KHOI[1:]]
	tieu_de += [nhan for nhan, _f, _kieu, _rong in COT_DOI_CHIEU]
	for c, t in enumerate(tieu_de, 1):
		o = ws.cell(4, c, t)
		o.font, o.alignment, o.border, o.fill = chu_tieu_de_cot, giua, vien, _nen(mau_cot(c))
	ws.row_dimensions[4].height = 31.5

	for r, dong in enumerate(data, 5):
		for c, (key, loai) in enumerate(truong, 1):
			gia_tri = dong.get(key)
			if loai in ("thue_suat", "chenh_lech_pt") and gia_tri is not None:
				gia_tri = flt(gia_tri) / 100
			elif loai == "ngay_hd" and gia_tri:
				gia_tri = getdate(gia_tri)
			o = ws.cell(r, c, gia_tri)
			o.border, o.alignment, o.font = vien, tren, chu
			if loai in _DINH_DANG:
				o.number_format = _DINH_DANG[loai]
			if dong.get("la_tong"):
				o.font, o.fill = chu_tong, _nen(_MAU_TONG)
			elif key == "so_hd_ban":
				o.fill = _nen(_MAU_SO_HD_RA)
			elif key == "trang_thai" and gia_tri != KHOP:
				o.font = _DO
		if dong.get("la_tong"):
			ws.cell(r, 3 + _TRUONG_KHOI.index("ten_hang"), "TỔNG CỘNG")

	_gop_manh(ws, data, truong)
	if data:
		_tom_tat(ws, 5 + len(data) + 1, data[-1], chu, chu_tong)

	for c, (_key, loai) in enumerate(truong, 1):
		ws.column_dimensions[get_column_letter(c)].width = _RONG.get(loai, 15)
	# Cố định tới hết cột Tên hàng bán: cuộn sang phải vẫn thấy hàng hoá.
	ws.freeze_panes = ws.cell(5, 3 + _TRUONG_KHOI.index("ten_hang") + 1)

	ra = BytesIO()
	wb.save(ra)
	return ra.getvalue()


def _gop_manh(ws, data, truong):
	"""Gộp ô theo chiều dọc cho các dòng liền nhau cùng STT (mảnh của một dòng bán)."""
	cot_gop = [c for c, (key, _loai) in enumerate(truong, 1) if key in _GOP]
	dau = 0
	for i in range(1, len(data) + 1):
		cung = i < len(data) and not data[i].get("la_tong") and data[i].get("stt") == data[dau].get("stt")
		if cung:
			continue
		if i - dau > 1:
			for c in cot_gop:
				ws.merge_cells(start_row=5 + dau, start_column=c, end_row=5 + i - 1, end_column=c)
				ws.cell(5 + dau, c).alignment = Alignment(vertical="top", wrap_text=c != 1)
		dau = i


def _tom_tat(ws, dong, tong, chu, chu_tong):
	tt = tong.get("tom_tat") or {}
	dem = tt.get("dem") or {}
	ty_le = tong.get("ty_le_gia")
	cac_dong = [
		("Doanh thu có đầu vào khớp", tt.get("doanh_thu_co_dau_vao"), _TIEN_AM_DO),
		("Doanh thu thiếu đầu vào", tt.get("doanh_thu_thieu_dau_vao"), _TIEN_AM_DO),
		("Doanh thu ghép theo mã hàng (thiếu liên kết)", tt.get("doanh_thu_theo_ma"), _TIEN_AM_DO),
		("Số dòng ghép theo mã hàng", tt.get("so_dong_theo_ma"), None),
		("Chênh lệch thành tiền (dòng đã ghép)", tong.get("chenh_lech_tien"), _TIEN_AM_DO),
		("Tỷ lệ giá bình quân (dòng đã ghép)", ty_le, "0.000000"),
		("Số HĐ bán / mua", f"{tt.get('so_hd_ban', 0)} / {tt.get('so_hd_mua', 0)}", None),
		("Số dòng lệch (đã ghép)", tt.get("so_dong_lech"), None),
		("Số dòng theo trạng thái", ", ".join(f"{k}: {v}" for k, v in dem.items() if v), None),
	]
	ws.cell(dong, 5, "TÓM TẮT").font = chu_tong
	for i, (nhan, gia_tri, dinh_dang) in enumerate(cac_dong, 1):
		ws.cell(dong + i, 5, nhan).font = chu
		o = ws.cell(dong + i, 9, gia_tri)
		o.font = chu_tong
		if dinh_dang:
			o.number_format = dinh_dang
