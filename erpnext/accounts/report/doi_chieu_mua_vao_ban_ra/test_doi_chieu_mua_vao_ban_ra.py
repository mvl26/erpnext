# Copyright (c) 2026, Công ty TNHH Miyano Việt Nam

"""Báo cáo Đối chiếu mua vào - bán ra trên chứng từ thật: SO → PO → PI và SO → SI."""

import json
from io import BytesIO

import frappe
import openpyxl
from frappe.tests.utils import FrappeTestCase
from frappe.utils import add_months, get_first_day, getdate, nowdate

from erpnext.accounts.report.doi_chieu_mua_vao_ban_ra.doi_chieu_mua_vao_ban_ra import execute, khoang_ky
from erpnext.buying.doctype.purchase_order.purchase_order import make_purchase_invoice
from erpnext.einvoice.tests.test_fixtures import ensure_item, ensure_uom
from erpnext.selling.doctype.sales_order.sales_order import make_purchase_order, make_sales_invoice

COMPANY = "Miyano"
WAREHOUSE = "Stores - M"


def _tk_vat():
	return frappe.db.get_value("Account", {"company": COMPANY, "account_type": "Tax", "is_group": 0}, "name")


def _vat(rate=5, mua=False):
	row = {"charge_type": "On Net Total", "account_head": _tk_vat(), "rate": rate, "description": "VAT"}
	if mua:
		row.update({"category": "Total", "add_deduct_tax": "Add"})
	return [row]


def _luu(doc, submit=True):
	doc.flags.ignore_permissions = True
	doc.insert()
	if submit:
		doc.submit()
	return doc


def _doi_tac(doctype):
	ten = f"_Test DCMB {doctype} {frappe.generate_hash(length=6)}"
	if doctype == "Customer":
		doc = {"customer_name": ten, "customer_group": "All Customer Groups", "territory": "All Territories"}
	else:
		doc = {"supplier_name": ten, "supplier_group": "All Supplier Groups"}
	return frappe.get_doc({"doctype": doctype, **doc}).insert(ignore_permissions=True).name


def _so(customer, qty, rate, ngay, item=None, ma_them=()):
	return _luu(
		frappe.get_doc(
			{
				"doctype": "Sales Order",
				"company": COMPANY,
				"customer": customer,
				"transaction_date": ngay,
				"delivery_date": ngay,
				"currency": "VND",
				"conversion_rate": 1,
				"items": [
					{
						"item_code": ma,
						"qty": qty,
						"rate": rate,
						"uom": ensure_uom(),
						"conversion_factor": 1,
						"warehouse": WAREHOUSE,
						"delivery_date": ngay,
					}
					for ma in (item or ensure_item(), *ma_them)
				],
				"taxes": _vat(),
			}
		)
	)


def _po(so, supplier, rate, ngay, item=None, bo_dong_so=False, them_tay=()):
	"""PO lập từ SO; `bo_dong_so` xoá `sales_order_item` (dòng PO bị sửa), giữ `sales_order`;
	`them_tay` thêm dòng hàng tay (không `sales_order`) vào cùng PO."""
	po = make_purchase_order(so.name, [{"item_code": item or ensure_item()}])
	for ma, qty in them_tay:
		po.append(
			"items",
			{
				"item_code": ma,
				"qty": qty,
				"uom": ensure_uom(),
				"conversion_factor": 1,
				"warehouse": WAREHOUSE,
			},
		)
	po.supplier = supplier
	po.transaction_date = ngay
	po.schedule_date = ngay
	po.taxes = []
	for r in _vat(mua=True):
		po.append("taxes", r)
	for it in po.items:
		it.rate = rate
		it.schedule_date = ngay
		if bo_dong_so:
			it.sales_order_item = None
	return _luu(po)


def _po_tay(supplier, qty, rate, ngay, item=None):
	"""PO lập tay, không từ SO: dòng PO không có `sales_order_item`."""
	return _luu(
		frappe.get_doc(
			{
				"doctype": "Purchase Order",
				"company": COMPANY,
				"supplier": supplier,
				"transaction_date": ngay,
				"schedule_date": ngay,
				"currency": "VND",
				"conversion_rate": 1,
				"items": [
					{
						"item_code": item or ensure_item(),
						"qty": qty,
						"rate": rate,
						"uom": ensure_uom(),
						"conversion_factor": 1,
						"warehouse": WAREHOUSE,
						"schedule_date": ngay,
					}
				],
				"taxes": _vat(mua=True),
			}
		)
	)


def _pi(po, qty, bill_no, ngay):
	pi = make_purchase_invoice(po.name)
	pi.bill_no = bill_no
	pi.bill_date = ngay
	pi.posting_date = ngay
	pi.set_posting_time = 1
	pi.update_stock = 0
	pi.items[0].qty = qty
	return _luu(pi)


def _si_tu_so(so, ngay, qty=None):
	si = make_sales_invoice(so.name)
	si.posting_date = ngay
	si.set_posting_time = 1
	si.update_stock = 0
	if qty is not None:
		si.items[0].qty = qty
	return _luu(si)


def _si_thang(customer, dong, ngay, item=None):
	return _luu(
		frappe.get_doc(
			{
				"doctype": "Sales Invoice",
				"company": COMPANY,
				"customer": customer,
				"posting_date": ngay,
				"set_posting_time": 1,
				"currency": "VND",
				"conversion_rate": 1,
				"items": [
					{
						"item_code": item or ensure_item(),
						"qty": q,
						"rate": r,
						"uom": ensure_uom(),
						"conversion_factor": 1,
					}
					for q, r in dong
				],
				"taxes": _vat(),
			}
		)
	)


def _du_lieu(**filters):
	_cot, data = execute({"company": COMPANY, "thang": getdate(nowdate()).strftime("%m-%Y"), **filters})
	return [r for r in data if not r.get("la_tong")], next((r for r in data if r.get("la_tong")), None)


class TestDoiChieuMuaVaoBanRa(FrappeTestCase):
	def setUp(self):
		self.hom_nay = nowdate()
		self.customer = _doi_tac("Customer")
		self.supplier = _doi_tac("Supplier")
		self.ten_ncc = frappe.db.get_value("Supplier", self.supplier, "supplier_name")

	def _chuoi(
		self, qty=10, rate_ban=101000, rate_mua=100000, pi_specs=(("00011069", 10),), ngay_pi=None, item=None
	):
		so = _so(self.customer, qty, rate_ban, self.hom_nay, item)
		po = _po(so, self.supplier, rate_mua, self.hom_nay, item)
		for bill_no, q in pi_specs:
			_pi(po, q, bill_no, ngay_pi or self.hom_nay)
		return so

	def _ma(self):
		"""Mã hàng riêng cho test: PI dự phòng ghép theo mã cho mọi khách, nên dòng bán của
		các test khác trong lớp (cùng mã chung, chưa rollback) sẽ tiêu mất PI của test này."""
		ma = f"_TEST-DCMB-{frappe.generate_hash(length=6)}"
		moi = frappe.copy_doc(frappe.get_doc("Item", ensure_item()))
		moi.update({"item_code": ma, "item_name": ma})
		return moi.insert(ignore_permissions=True).name

	def test_ghep_mot_dong(self):
		si = _si_tu_so(self._chuoi(), self.hom_nay)
		dong, tong = _du_lieu(customer=self.customer)
		self.assertEqual(len(dong), 1)
		r = dong[0]
		self.assertEqual(r["stt"], 1)
		self.assertEqual(r["so_hd_ban"], "")
		self.assertEqual(r["sales_invoice"], si.name)
		self.assertIn("Chưa có số HĐĐT", r["ghi_chu"])
		self.assertEqual(r["trang_thai"], "Khớp")
		self.assertEqual(r["ty_le_gia"], 1.01)
		self.assertEqual(r["chenh_lech_tien"], 10000)
		self.assertEqual(r["thanh_tien_ban"], 1010000)
		self.assertEqual(r["tien_thue_ban"], 50500)
		self.assertEqual(r["thue_suat_ban"], 5)
		self.assertEqual(r["ncc"], self.ten_ncc)
		self.assertEqual(r["so_hd_mua"], "00011069")
		self.assertEqual(r["thanh_tien_mua"], 1000000)
		self.assertEqual(r["tien_thue_mua"], 50000)
		self.assertEqual(tong["tong_ban"], 1060500)
		self.assertEqual(tong["tong_mua"], 1050000)

	def test_khong_ghep_van_hien(self):
		_si_thang(self.customer, [(2, 50000)], self.hom_nay)
		dong, _tong = _du_lieu(customer=self.customer)
		self.assertEqual(len(dong), 1)
		self.assertFalse(dong[0].get("ncc"))
		self.assertIsNone(dong[0].get("thanh_tien_mua"))
		self.assertEqual(dong[0]["trang_thai"], "Thiếu đầu vào")

	def test_ghep_theo_ma_khi_po_khong_lap_tu_so(self):
		ma = self._ma()
		so = _so(self.customer, 10, 101000, self.hom_nay, ma)
		_si_tu_so(so, self.hom_nay)
		_pi(_po_tay(self.supplier, 10, 100000, self.hom_nay, ma), 10, "00022001", self.hom_nay)
		dong, tong = _du_lieu(customer=self.customer)
		self.assertEqual(len(dong), 1)
		r = dong[0]
		self.assertEqual((r["trang_thai"], r["so_hd_mua"], r["theo_ma"]), ("Khớp", "00022001", 1))
		self.assertIn("Ghép theo mã hàng – thiếu liên kết SO→PO", r["ghi_chu"])
		self.assertEqual(tong["tom_tat"]["so_dong_theo_ma"], 1)

	def test_ghep_theo_ma_hd_ban_khong_qua_so(self):
		ma = self._ma()
		_si_thang(self.customer, [(2, 50500)], self.hom_nay, ma)
		_pi(_po_tay(self.supplier, 2, 50000, self.hom_nay, ma), 2, "00022002", self.hom_nay)
		dong, _tong = _du_lieu(customer=self.customer)
		self.assertEqual(dong[0]["trang_thai"], "Khớp")
		self.assertIn("Ghép theo mã hàng – HĐ bán không qua SO", dong[0]["ghi_chu"])

	def test_pi_cua_so_da_xuat_hd_khong_ghep_theo_ma(self):
		# PI thuộc chuỗi của một SO đã xuất HĐ bán (khách khác, tháng sau) không được đem cho
		# dòng bán này: ngày HĐ bán của SO kia không quan trọng.
		ma = self._ma()
		khach_khac = _doi_tac("Customer")
		so = _so(khach_khac, 10, 101000, self.hom_nay, ma)
		_pi(_po(so, self.supplier, 100000, self.hom_nay, ma), 10, "00022005", self.hom_nay)
		si = make_sales_invoice(so.name)
		si.posting_date = si.due_date = add_months(self.hom_nay, 1)
		si.set_posting_time = 1
		si.update_stock = 0
		si.payment_schedule = []
		_luu(si)
		_si_thang(self.customer, [(2, 101000)], self.hom_nay, ma)
		dong, _tong = _du_lieu(customer=self.customer)
		self.assertEqual(dong[0]["trang_thai"], "Thiếu đầu vào")

	def test_pi_cua_so_chua_xuat_hd_duoc_ghep_theo_ma(self):
		# SO lập trùng: SO đầu sinh PO → PI nhưng không xuất HĐ; HĐ bán xuất từ SO thứ hai.
		ma = self._ma()
		self._chuoi(item=ma, pi_specs=(("00022004", 10),))
		_si_tu_so(_so(self.customer, 10, 101000, self.hom_nay, ma), self.hom_nay)
		dong, _tong = _du_lieu(customer=self.customer)
		self.assertEqual(len(dong), 1)
		self.assertEqual((dong[0]["trang_thai"], dong[0]["so_hd_mua"]), ("Khớp", "00022004"))
		self.assertIn("Ghép theo mã hàng – thiếu liên kết SO→PO", dong[0]["ghi_chu"])

	def test_ghep_trong_cap_so_po_khi_dong_po_mat_lien_ket(self):
		ma = self._ma()
		so = _so(self.customer, 10, 101000, self.hom_nay, ma)
		po = _po(so, self.supplier, 100000, self.hom_nay, ma, bo_dong_so=True)
		self.assertEqual((po.items[0].sales_order, po.items[0].sales_order_item), (so.name, None))
		_pi(po, 10, "00022006", self.hom_nay)
		_si_tu_so(so, self.hom_nay)
		dong, tong = _du_lieu(customer=self.customer)
		self.assertEqual(len(dong), 1)
		r = dong[0]
		self.assertEqual((r["trang_thai"], r["so_hd_mua"], r.get("theo_cap")), ("Khớp", "00022006", 1))
		self.assertIsNone(r.get("theo_ma"))
		self.assertIn(f"Ghép trong cặp SO→PO ({so.name} → {po.name})", r["ghi_chu"])
		self.assertNotIn("thiếu liên kết", r["ghi_chu"])
		self.assertEqual((tong["tom_tat"]["so_dong_theo_cap"], tong["tom_tat"]["so_dong_theo_ma"]), (1, 0))

	def test_ghep_trong_cap_khi_dong_po_them_tay(self):
		# Mẫu prod (PO-00078): PO lập từ SO cho mã A, mã B thêm tay vào cùng PO.
		ma_a, ma_b = self._ma(), self._ma()
		so = _so(self.customer, 10, 101000, self.hom_nay, ma_a, ma_them=(ma_b,))
		po = _po(so, self.supplier, 100000, self.hom_nay, ma_a, them_tay=((ma_b, 10),))
		self.assertEqual([it.sales_order for it in po.items], [so.name, None])
		_pi(po, 10, "00022007", self.hom_nay)
		_si_tu_so(so, self.hom_nay)
		dong, tong = _du_lieu(customer=self.customer)
		self.assertEqual([r["trang_thai"] for r in dong], ["Khớp", "Khớp"])
		self.assertEqual([r.get("theo_cap") for r in dong], [None, 1])
		self.assertIn(f"Ghép trong cặp SO→PO ({so.name} → {po.name})", dong[1]["ghi_chu"])
		self.assertEqual((tong["tom_tat"]["so_dong_theo_cap"], tong["tom_tat"]["so_dong_theo_ma"]), (1, 0))

	def test_pi_du_phong_khong_ghep_lai_o_ky_sau(self):
		thang_truoc = add_months(get_first_day(self.hom_nay), -1)
		ma = self._ma()
		_pi(_po_tay(self.supplier, 3, 100000, thang_truoc, ma), 3, "00022003", thang_truoc)
		_si_thang(self.customer, [(2, 101000)], thang_truoc, ma)
		_si_thang(self.customer, [(2, 101000)], self.hom_nay, ma)
		dong, _tong = _du_lieu(customer=self.customer)
		self.assertEqual([(r["sl_ban"], r.get("so_hd_mua")) for r in dong], [(1, "00022003"), (1, None)])

	def test_mot_dong_ban_hai_hd_mua(self):
		_si_tu_so(self._chuoi(pi_specs=(("A", 6), ("B", 4))), self.hom_nay)
		dong, tong = _du_lieu(customer=self.customer)
		self.assertEqual([r["so_hd_mua"] for r in dong], ["A", "B"])
		self.assertEqual([r["sl_ban"] for r in dong], [6, 4])
		self.assertEqual([r["stt"] for r in dong], [1, 1])
		self.assertEqual([r["manh"] for r in dong], ["1/2", "2/2"])
		self.assertEqual(tong["thanh_tien_ban"], 1010000)
		self.assertEqual(tong["thanh_tien_mua"], 1000000)

	def test_loc_ncc(self):
		_si_tu_so(self._chuoi(), self.hom_nay)
		_si_thang(self.customer, [(2, 50000)], self.hom_nay)
		self.assertEqual(len(_du_lieu(customer=self.customer)[0]), 2)
		dong, _tong = _du_lieu(customer=self.customer, supplier=self.supplier)
		self.assertEqual(len(dong), 1)
		self.assertEqual(dong[0]["ncc"], self.ten_ncc)

	def test_tong_khop_hai_dong_cung_ma(self):
		si = _si_thang(self.customer, [(1, 1611943), (1, 555905)], self.hom_nay)
		_dong, tong = _du_lieu(customer=self.customer)
		self.assertEqual(tong["thanh_tien_ban"], si.base_net_total)
		self.assertEqual(tong["tien_thue_ban"], si.base_total_taxes_and_charges)

	def test_ky_loc_theo_ngay_ban(self):
		thang_truoc = add_months(get_first_day(self.hom_nay), -1)
		_si_thang(self.customer, [(1, 70000)], thang_truoc)
		_si_tu_so(self._chuoi(ngay_pi=thang_truoc), self.hom_nay)
		dong, _tong = _du_lieu(customer=self.customer)
		self.assertEqual(len(dong), 1)
		self.assertEqual(getdate(dong[0]["ngay_hd_mua"]), getdate(thang_truoc))

	def test_ban_vuot_mua_va_loc_trang_thai(self):
		_si_tu_so(self._chuoi(pi_specs=(("A", 6),)), self.hom_nay)
		_si_thang(self.customer, [(2, 50000)], self.hom_nay)
		dong, tong = _du_lieu(customer=self.customer)
		self.assertEqual([r["trang_thai"] for r in dong], ["Khớp", "Thiếu đầu vào", "Thiếu đầu vào"])
		self.assertEqual(tong["tom_tat"]["doanh_thu_co_dau_vao"], 606000)
		self.assertEqual(tong["tom_tat"]["doanh_thu_thieu_dau_vao"], 404000 + 100000)
		self.assertEqual(tong["ty_le_gia"], 1.01)

		dong, tong = _du_lieu(customer=self.customer, trang_thai="Thiếu đầu vào")
		self.assertEqual(len(dong), 2)
		self.assertEqual(dong[0]["manh"], "2/2")
		self.assertEqual(tong["thanh_tien_ban"], 504000)
		self.assertIsNone(tong["ty_le_gia"])
		self.assertEqual(len(_du_lieu(customer=self.customer, trang_thai="Khớp")[0]), 1)

	def test_lech_gia_theo_ty_le_hop_dong(self):
		_si_tu_so(self._chuoi(rate_ban=102000), self.hom_nay)
		self.assertEqual(_du_lieu(customer=self.customer)[0][0]["trang_thai"], "Lệch giá")
		self.assertEqual(_du_lieu(customer=self.customer, ty_le_hd=1.02)[0][0]["trang_thai"], "Khớp")

	def test_thieu_so_hd_ncc_va_hd_mua_ngoai_ky(self):
		thang_truoc = add_months(get_first_day(self.hom_nay), -1)
		_si_tu_so(self._chuoi(pi_specs=((None, 10),), ngay_pi=thang_truoc), self.hom_nay)
		r = _du_lieu(customer=self.customer)[0][0]
		self.assertEqual(r["so_hd_mua"], "")
		self.assertTrue(r["purchase_invoice"])
		self.assertIn("Thiếu số HĐ NCC", r["ghi_chu"])
		self.assertIn("ngoài kỳ", r["ghi_chu"])

	def test_tra_hang(self):
		from erpnext.controllers.sales_and_purchase_return import make_return_doc

		si = _si_tu_so(self._chuoi(), self.hom_nay)
		tra = make_return_doc("Sales Invoice", si.name)
		tra.items[0].qty = -2
		_luu(tra)
		dong, tong = _du_lieu(customer=self.customer)
		r = next(d for d in dong if d["tra_hang"])
		self.assertEqual(r["trang_thai"], "Trả hàng")
		self.assertEqual(r["tra_cho"], si.name)
		self.assertIn(f"Trả lại HĐ {si.name}", r["ghi_chu"])
		self.assertEqual(r["sl_ban"], -2)
		self.assertEqual(tong["tom_tat"]["dem"]["Trả hàng"], 1)

	def test_ky_trong(self):
		dong, tong = _du_lieu(customer=self.customer)
		self.assertEqual(dong, [])
		self.assertIsNone(tong)

	def test_khoang_ky(self):
		f = frappe._dict
		self.assertEqual(
			khoang_ky(f(thang="10-2026")), (getdate("2026-10-01"), getdate("2026-10-31"), "THÁNG 10.2026")
		)
		self.assertEqual(
			khoang_ky(f(thang="10-2026", from_date="2026-10-01", to_date="2026-10-20"))[2],
			"TỪ 01/10/2026 ĐẾN 20/10/2026",
		)
		self.assertEqual(khoang_ky(f(from_date="2026-02-01", to_date="2026-02-28"))[2], "THÁNG 02.2026")
		for sai in ("10/2026", "2026-10", "13-2026"):
			with self.assertRaises(frappe.ValidationError):
				khoang_ky(f(thang=sai))


class TestXuatExcel(FrappeTestCase):
	def setUp(self):
		self.customer = _doi_tac("Customer")
		self.filters = {
			"company": COMPANY,
			"customer": self.customer,
			"thang": getdate(nowdate()).strftime("%m-%Y"),
		}

	def tearDown(self):
		frappe.set_user("Administrator")

	def _doc(self):
		from erpnext.accounts.report.doi_chieu_mua_vao_ban_ra.xuat_excel import dung_workbook

		return openpyxl.load_workbook(BytesIO(dung_workbook(self.filters))).active

	def test_xuat_excel_bo_cuc(self):
		supplier = _doi_tac("Supplier")
		so = _so(self.customer, 10, 101000, nowdate())
		_pi(_po(so, supplier, 100000, nowdate()), 10, "00011069", nowdate())
		_si_tu_so(so, nowdate())
		ws = self._doc()
		self.assertIn("THÁNG", ws["A1"].value)
		self.assertTrue(ws["B3"].value.startswith("KHÁCH HÀNG"))
		self.assertEqual(ws["M3"].value, "NHÀ CUNG CẤP")
		self.assertEqual(ws["C4"].value, "Số HĐ đầu ra")
		self.assertEqual(ws["N4"].value, "Số HĐ đầu vào")
		self.assertEqual(ws["I5"].value, 1010000)
		self.assertEqual(ws["T5"].value, 1000000)
		self.assertEqual(ws["N5"].value, "00011069")
		self.assertEqual(ws["J5"].value, 0.05)
		self.assertEqual(ws["J5"].number_format, "0%")
		self.assertEqual(ws["E6"].value, "TỔNG CỘNG")
		self.assertEqual(ws["L6"].value, 1060500)
		self.assertEqual(ws["W6"].value, 1050000)
		self.assertIn("B3:L3", {str(r) for r in ws.merged_cells.ranges})
		self.assertIn("M3:W3", {str(r) for r in ws.merged_cells.ranges})
		self.assertEqual(ws["X3"].value, "ĐỐI CHIẾU")
		self.assertEqual(ws["X4"].value, "Tỷ lệ giá bán / mua")
		self.assertEqual(ws["X5"].value, 1.01)
		self.assertEqual(ws["AA5"].value, "Khớp")
		self.assertEqual(ws["Z6"].value, 10000)
		self.assertEqual(ws.freeze_panes, "F5")
		self.assertEqual(ws["E8"].value, "TÓM TẮT")

	def test_xuat_excel_gop_o_manh(self):
		supplier = _doi_tac("Supplier")
		so = _so(self.customer, 10, 101000, nowdate())
		po = _po(so, supplier, 100000, nowdate())
		_pi(po, 6, "A", nowdate())
		_pi(po, 4, "B", nowdate())
		_si_tu_so(so, nowdate())
		ws = self._doc()
		gop = {str(r) for r in ws.merged_cells.ranges}
		self.assertTrue({"A5:A6", "C5:C6", "E5:E6", "H5:H6"} <= gop)
		self.assertNotIn("G5:G6", gop)
		self.assertEqual((ws["G5"].value, ws["G6"].value), (6, 4))

	def test_xuat_excel_ky_trong(self):
		ws = self._doc()
		self.assertIn("THÁNG", ws["A1"].value)
		self.assertIsNone(ws["A5"].value)

	def _dang_nhap(self, role):
		email = f"_test_dcmb_{frappe.generate_hash(length=6)}@example.com"
		user = frappe.get_doc(
			{"doctype": "User", "email": email, "first_name": "DCMB", "send_welcome_email": 0}
		).insert(ignore_permissions=True)
		user.add_roles(role)
		frappe.set_user(email)

	def test_xuat_excel_chan_nguoi_khong_quyen(self):
		from erpnext.accounts.report.doi_chieu_mua_vao_ban_ra.xuat_excel import xuat

		self._dang_nhap("Kế toán HĐĐT" if frappe.db.exists("Role", "Kế toán HĐĐT") else "Sales User")
		with self.assertRaises(frappe.PermissionError):
			xuat(json.dumps(self.filters))

	def test_xuat_excel_cho_accounts_user(self):
		from erpnext.accounts.report.doi_chieu_mua_vao_ban_ra.xuat_excel import xuat

		self._dang_nhap("Accounts User")
		xuat(json.dumps(self.filters))
		self.assertEqual(frappe.response.type, "binary")
		self.assertTrue(frappe.response.filename.endswith(".xlsx"))
		self.assertTrue(frappe.response.filecontent.startswith(b"PK"))

	def test_du_lieu_trang(self):
		from erpnext.accounts.report.doi_chieu_mua_vao_ban_ra.xuat_excel import du_lieu_trang

		_si_thang(self.customer, [(2, 50000)], nowdate())
		kq = du_lieu_trang(json.dumps(self.filters))
		self.assertTrue(kq["tieu_de"].startswith("TỔNG HỢP ĐẦU VÀO - RA CỦA KHÁCH HÀNG _TEST DCMB"))
		self.assertEqual(len(kq["data"]), 2)
		self.assertEqual(kq["data"][-1]["thanh_tien_ban"], 100000)
		self.assertEqual(du_lieu_trang(json.dumps({"company": COMPANY}))["data"], [])

	def test_du_lieu_trang_chan_nguoi_khong_quyen(self):
		from erpnext.accounts.report.doi_chieu_mua_vao_ban_ra.xuat_excel import du_lieu_trang

		self._dang_nhap("Sales User")
		with self.assertRaises(frappe.PermissionError):
			du_lieu_trang(json.dumps(self.filters))
